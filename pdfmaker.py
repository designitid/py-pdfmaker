from fpdf import FPDF
from docx import Document

# Text content - Paste your  text here

content = """
Lorem ipsum dolor sit amet, consectetur adipiscing elit. 
Sed sit amet nulla auctor, vestibulum magna sed, convallis ex.

"""

# Create PDF
pdf = FPDF()
pdf.add_page()
pdf.set_font("Arial", size=12)
for line in content.strip().split('\n'):
    pdf.cell(200, 10, txt=line, ln=True)

pdf_output_path = "/mnt/data/jaringan_komputer.pdf"
pdf.output(pdf_output_path)

# Create Word document
doc = Document()
for line in content.strip().split('\n'):
    doc.add_paragraph(line)

word_output_path = "/mnt/data/jaringan_komputer.docx"
doc.save(word_output_path)

pdf_output_path, word_output_path
